# регрессия по исправлениям ревью (только реальные сценарии, без синтетики). Запуск: pytest -q tests
import io
import pytest
from app import routing, store
from app.nlp import extractor
from app.nlp.extractor import extract

store.init()


def ids(text):
    return {f.id for f in extract(text, None).findings}


def reasons(text, fid):
    return [s.reason for s in extract(text, None).suppressed if s.id == fid]


# словарь органов и нечёткое сравнение
def test_common_word_is_not_fuzzy_organ():
    # «простой» — обычное слово, а не опечатка «простата»: киста яичника находится
    assert "ovarian_mass" in ids("Эхографические признаки простой кисты правого яичника.")


def test_abbreviation_is_exact_only():
    # «обеих» (начальная форма «оба») — не артерия ОБА
    assert "thyroid_nodule" in ids("Заключение\nУзлы в обеих долях щитовидной железы.")


def test_lymph_nodes_are_not_thyroid_nodule():
    t = "УЗИ щитовидной железы\nЩитовидная железа: структура однородная.\nРегионарные лимфатические узлы не изменены."
    assert "thyroid_nodule" not in ids(t)


def test_typo_with_transposition_found():
    assert "gallbladder_polyp" in ids("Полип желчного пзуыря 13 мм.")


def test_section_not_broken_by_reference_organ_line():
    t = "ПРАВЫЙ ЯИЧНИК\nРазмеры: 24х13х16 мм\nФаллопиева труба: не визуализируется\nПараовариально лоцируется анэхогенное образование 12х11мм"
    assert "ovarian_mass" in ids(t)


# контекстные исключения
def test_aortic_thrombus_is_not_leg_vein_thrombosis():
    t = "Инфраренальный отдел брюшной аорты расширен до 34 мм, пристеночные тромботические массы."
    assert "venous_thrombosis" not in ids(t)
    assert "venous_thrombosis" in ids("Окклюзивный тромбоз подколенной вены справа.")


def test_admin_exclusion_applies_and_can_be_removed():
    t = "Пупочная грыжа, грыжевые ворота 15 мм."
    assert "hernia" in ids(t)
    extractor.set_user_exclusions([{"finding_id": "hernia", "phrase": "пупочный", "scope": "sentence"}])
    try:
        assert "hernia" not in ids(t)
        assert any("администратора" in r for r in reasons(t, "hernia"))
    finally:
        extractor.set_user_exclusions([])
    assert "hernia" in ids(t)


# BI-RADS, размер из описания, поглощение, экстренные
def test_birads_both_sides_one_record():
    t = "Молочные железы: образования не выявлены.\nЗаключение: BI-RADS справа 2 , слева 2"
    assert "breast_birads" not in ids(t)


def test_single_birads_for_whole_exam():
    t = "В обеих молочных железах анэхогенные округлые образования до 8 мм.\nЗАКЛЮЧЕНИЕ: кистозная мастопатия, категория Birads 2."
    assert "breast_birads" not in ids(t)


def test_gallbladder_polyp_size_taken_from_description():
    t = ("ЖЕЛЧНЫЙ ПУЗЫРЬ: пристеночное образование повышенной эхогенности, без акустической тени, размером 3,6*3,5 мм -полип\n"
         "ЗАКЛЮЧЕНИЕ: Образование в желчном пузыре- полип (без динамики)")
    assert "gallbladder_polyp" not in ids(t)


def test_general_endometrial_pathology_subsumed_by_polyp():
    found = ids("ЭХО- признаки патологии эндометрия ( полип ), аденомиоза.")
    assert "endometrial_polyp" in found and "endometrial_pathology" not in found


def test_ectopic_pregnancy_is_urgent():
    f = next(f for f in extract("Нельзя исключить внематочную беременность справа.", None).findings if f.id == "ectopic_pregnancy")
    assert f.urgent


# .docx
def test_docx_text_extraction():
    import docx
    from app.main import _docx_text
    d = docx.Document()
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Заключение:"
    t.cell(0, 1).text = "Полип желчного пузыря 12 мм."
    buf = io.BytesIO(); d.save(buf)
    text = _docx_text(buf.getvalue())
    assert "Полип желчного пузыря 12 мм." in text
    assert "gallbladder_polyp" in ids(text)


# автоотправка без проверки
def _route(pid):
    return store.q("SELECT * FROM routes WHERE protocol_id=?", (pid,))


def test_auto_send_off_by_default():
    store.set_setting("auto_send", {"mode": "off"})
    routing.handle_event({"event_id": "V5-1", "type": "protocol.signed", "protocol_id": "V5-P1", "patient_id": "v5-1",
                          "text": "Полип желчного пузыря 12 мм."})
    assert _route("V5-P1")[0]["status"] == "pending_verification"


def test_auto_send_all_notifies_patient_immediately():
    store.set_setting("auto_send", {"mode": "all", "skip_needs_review": True})
    routing.handle_event({"event_id": "V5-2", "type": "protocol.signed", "protocol_id": "V5-P2", "patient_id": "v5-2",
                          "text": "Полип желчного пузыря 12 мм."})
    r = _route("V5-P2")[0]
    assert r["status"] == "notified"
    assert store.q("SELECT COUNT(*) AS n FROM notifications WHERE route_id=? AND kind='patient'", (r["id"],))[0]["n"] >= 1


def test_auto_send_skips_findings_that_need_review():
    store.set_setting("auto_send", {"mode": "all", "skip_needs_review": True})
    routing.handle_event({"event_id": "V5-3", "type": "protocol.signed", "protocol_id": "V5-P3", "patient_id": "v5-3",
                          "text": "Полип желчного пузыря."})          # нет размера -> пометка «проверить»
    assert _route("V5-P3")[0]["status"] == "pending_verification"


def test_auto_send_never_for_urgent():
    store.set_setting("auto_send", {"mode": "all", "skip_needs_review": False})
    routing.handle_event({"event_id": "V5-4", "type": "protocol.signed", "protocol_id": "V5-P4", "patient_id": "v5-4",
                          "text": "Окклюзивный тромбоз подколенной вены справа."})
    r = _route("V5-P4")[0]
    assert r["status"] == "escalated"
    assert store.q("SELECT COUNT(*) AS n FROM notifications WHERE route_id=? AND kind='patient'", (r["id"],))[0]["n"] == 0
    store.set_setting("auto_send", {"mode": "off"})


def test_auto_send_selected_routes_only():
    store.set_setting("auto_send", {"mode": "selected", "routes": ["urologist"], "skip_needs_review": True})
    routing.handle_event({"event_id": "V5-5", "type": "protocol.signed", "protocol_id": "V5-P5", "patient_id": "v5-5",
                          "text": "Полип желчного пузыря 12 мм. Конкремент правой почки 6 мм."})
    st = {r["route_key"]: r["status"] for r in _route("V5-P5")}
    assert st.get("urologist") == "notified" and st.get("surgeon") == "pending_verification"
    store.set_setting("auto_send", {"mode": "off"})
